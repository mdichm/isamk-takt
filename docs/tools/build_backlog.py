# -*- coding: utf-8 -*-
"""Génère Backlog_TAKT_tickets.csv, Backlog_TAKT.xlsx et la table de correspondance des anciens numéros."""
import csv, re, sys, json, os
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__)); DOCS = os.path.dirname(HERE)
exec(open(os.path.join(HERE, "tickets_data.py"), encoding="utf-8").read())

LOTS = OrderedDict([
 ("L0","Socle : monorepo, CI/CD, infrastructure sur site, authentification, sécurité, RBAC, audit, socle API et web"),
 ("L1","Années, semestres, référentiel, ressources, grille horaire, Ramadan, imports"),
 ("L2","Répartition pédagogique, charge paramétrable, snapshots, vue annuelle"),
 ("L3","Contraintes, versions, séances, conflict-engine, éditeur manuel, verrous, avertissements"),
 ("L4","Génération FET : port, profils, worker, périmètre, préparation, comparaison, score, adoption"),
 ("L5","Workflow, impact, notifications, circuit configurable, tableau de bord"),
 ("L6","Exports PDF et Excel"),
 ("L7","Clôture, durcissement, performance, exploitation, recette, mise en production"),
])
# attribution des identifiants et des epics
cnt = {l:0 for l in LOTS}; epics = OrderedDict(); epic_n = {l:0 for l in LOTS}
for tk in TICKETS:
    l = tk["lot"]; cnt[l] += 1; tk["id"] = f"{l}-{cnt[l]:02d}"
    ek = (l, tk["epic"])
    if ek not in epics:
        epic_n[l] += 1; epics[ek] = f"E{l[1]}.{epic_n[l]}"
    tk["epic_id"] = epics[ek]
by_key = {tk["key"]: tk for tk in TICKETS}
for tk in TICKETS:
    tk["deps_ids"] = [by_key[d]["id"] for d in tk["deps"]]

HEAD = ["ID","Lot","Epic ID","Epic","Type","Titre","Description","Critères d'acceptation","Réf. SF","Réf. ST","Profil","Priorité","Points","Dépendances","Statut","Sprint"]
def row(tk):
    return [tk["id"],tk["lot"],tk["epic_id"],tk["epic"],tk["typ"],tk["titre"],tk["desc"],
            "\n".join("- "+a for a in tk["ac"]),tk["sf"],tk["st"],tk["profil"],tk["prio"],tk["pts"],
            ", ".join(tk["deps_ids"]),tk["statut"],""]

# CSV
with open(os.path.join(DOCS,"Backlog_TAKT_tickets.csv"),"w",encoding="utf-8-sig",newline="") as f:
    w = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL); w.writerow(HEAD)
    for tk in TICKETS: w.writerow(row(tk))

# correspondance anciens numéros
with open(os.path.join(DOCS,"tools","correspondance_tickets.json"),"w",encoding="utf-8") as f:
    json.dump({tk["old"]:tk["id"] for tk in TICKETS if tk["old"]}, f, ensure_ascii=False, indent=1)

# XLSX
F = "Arial"
hfill = PatternFill("solid", fgColor="29363C"); hfont = Font(name=F, bold=True, color="FFFFFF", size=10)
body = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
thin = Side(style="thin", color="DDE5E4"); border = Border(left=thin,right=thin,top=thin,bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
wb = Workbook()

# Lisez-moi
ws = wb.active; ws.title = "Lisez-moi"
lines = [("TAKT — Backlog de référence, version 1.0",True),("Version 1.0 · 3 octobre 2026. Source des tickets : Backlog_TAKT_tickets.csv (séparateur « ; »).",False),("",False),
 ("Lecture",True),("Chaque ticket cite les exigences de la spécification fonctionnelle (codes EX-…, colonne « Réf. SF ») et les sections de la spécification technique (colonne « Réf. ST »).",False),
 ("Un ticket sans exigence est un ticket technique, de test ou un spike, justifié par la ST.",False),
 ("Les totaux des feuilles Lots et Epics sont calculés par formule depuis la feuille Tickets.",False),
 ("Les cellules Statut et Sprint de la feuille Tickets sont à renseigner par l'équipe.",False),
 ("La matrice de traçabilité est générée par tools/verifier_coherence.py (TRACABILITE_TAKT.md).",False)]
for i,(txt,b) in enumerate(lines,1):
    c = ws.cell(row=i,column=1,value=txt); c.font = Font(name=F,size=12 if i==1 else 10,bold=b)
ws.column_dimensions["A"].width = 130

# Tickets
wt = wb.create_sheet("Tickets")
for j,h in enumerate(HEAD,1):
    c = wt.cell(row=1,column=j,value=h); c.font=hfont; c.fill=hfill; c.alignment=Alignment(wrap_text=True,vertical="center"); c.border=border
for i,tk in enumerate(TICKETS,2):
    for j,v in enumerate(row(tk),1):
        c = wt.cell(row=i,column=j,value=v); c.font=body; c.alignment=wrap; c.border=border
n = len(TICKETS)+1
widths = [9,6,8,26,15,38,52,70,24,22,10,10,8,22,10,8]
for j,wd in enumerate(widths,1): wt.column_dimensions[get_column_letter(j)].width = wd
wt.freeze_panes = "C2"; wt.auto_filter.ref = f"A1:P{n}"

# Lots
wl = wb.create_sheet("Lots")
for j,h in enumerate(["Lot","Périmètre","Tickets","Points","Points terminés"],1):
    c = wl.cell(row=1,column=j,value=h); c.font=hfont; c.fill=hfill; c.border=border
for i,(l,d) in enumerate(LOTS.items(),2):
    wl.cell(row=i,column=1,value=l); wl.cell(row=i,column=2,value=d)
    wl.cell(row=i,column=3,value=f'=COUNTIF(Tickets!$B$2:$B${n},A{i})')
    wl.cell(row=i,column=4,value=f'=SUMIF(Tickets!$B$2:$B${n},A{i},Tickets!$M$2:$M${n})')
    wl.cell(row=i,column=5,value=f'=SUMIFS(Tickets!$M$2:$M${n},Tickets!$B$2:$B${n},A{i},Tickets!$O$2:$O${n},"Terminé")')
    for j in range(1,6):
        c = wl.cell(row=i,column=j); c.font=body; c.border=border; c.alignment=wrap
r = len(LOTS)+2
wl.cell(row=r,column=1,value="Total").font=bold
for col in "CDE":
    c = wl[f"{col}{r}"]; c.value=f"=SUM({col}2:{col}{r-1})"; c.font=bold
for j,wd in enumerate([8,90,10,10,16],1): wl.column_dimensions[get_column_letter(j)].width = wd

# Epics
we = wb.create_sheet("Epics")
for j,h in enumerate(["Epic ID","Lot","Epic","Tickets","Points"],1):
    c = we.cell(row=1,column=j,value=h); c.font=hfont; c.fill=hfill; c.border=border
for i,((l,name),eid) in enumerate(epics.items(),2):
    we.cell(row=i,column=1,value=eid); we.cell(row=i,column=2,value=l); we.cell(row=i,column=3,value=name)
    we.cell(row=i,column=4,value=f'=COUNTIF(Tickets!$C$2:$C${n},A{i})')
    we.cell(row=i,column=5,value=f'=SUMIF(Tickets!$C$2:$C${n},A{i},Tickets!$M$2:$M${n})')
    for j in range(1,6):
        c = we.cell(row=i,column=j); c.font=body; c.border=border
r = len(epics)+2
we.cell(row=r,column=1,value="Total").font=bold
for col in "DE":
    c = we[f"{col}{r}"]; c.value=f"=SUM({col}2:{col}{r-1})"; c.font=bold
for j,wd in enumerate([9,6,46,10,10],1): we.column_dimensions[get_column_letter(j)].width = wd
wb.save(os.path.join(DOCS,"Backlog_TAKT.xlsx"))
print(len(TICKETS),"tickets,",len(epics),"epics,",sum(t["pts"] for t in TICKETS),"points")
